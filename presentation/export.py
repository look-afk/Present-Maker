"""Editable PowerPoint export optimized for explanatory study-note slides."""
from pathlib import Path
from io import BytesIO
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import threading
_RENDER_LOCK = threading.Lock()

BG=RGBColor(12,18,32); PANEL=RGBColor(24,34,54); PANEL2=RGBColor(31,45,70)
ACCENT=RGBColor(102,183,255); WHITE=RGBColor(246,248,252); MUTED=RGBColor(174,188,210); GREEN=RGBColor(123,226,191)


def background(slide,color=BG):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color


def add_text(slide,value,x,y,w,h,size=18,color=WHITE,bold=False,font="Aptos",italic=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=shape.text_frame
    tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.06); tf.margin_right=Inches(.06); tf.margin_top=Inches(.04); tf.margin_bottom=Inches(.03)
    p=tf.paragraphs[0]; p.text=str(value); p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.italic=italic; p.font.color.rgb=color
    return shape


def rect(slide,x,y,w,h,color,rounded=True):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=color; shape.line.fill.background(); return shape


def add_image(slide,item,x,y,w,h):
    url=item.get("image_url")
    if not url:return False
    try:
        r=requests.get(url,timeout=10,headers={"User-Agent":"PresentationBot/1.0"}); r.raise_for_status()
        if len(r.content)>12_000_000:return False
        slide.shapes.add_picture(BytesIO(r.content), Inches(x), Inches(y), width=Inches(w), height=Inches(h)); return True
    except Exception:return False


def create_pptx(topic,lang,content,sources,output_path,settings=None):
    settings = settings or {}
    with _RENDER_LOCK:
        return _create_pptx_with_theme(topic, lang, content, sources, output_path, settings)


def _create_pptx_with_theme(topic,lang,content,sources,output_path,settings):
    global BG, PANEL, PANEL2, ACCENT, WHITE, MUTED, GREEN
    palettes = {
        "dark": (RGBColor(12,18,32), RGBColor(24,34,54), RGBColor(31,45,70), RGBColor(102,183,255), RGBColor(246,248,252), RGBColor(174,188,210), RGBColor(123,226,191)),
        "academic": (RGBColor(19,35,56), RGBColor(29,49,72), RGBColor(39,62,87), RGBColor(210,170,91), RGBColor(250,248,241), RGBColor(195,203,211), RGBColor(164,199,177)),
        "minimal": (RGBColor(248,249,251), RGBColor(255,255,255), RGBColor(235,239,244), RGBColor(39,91,140), RGBColor(30,38,48), RGBColor(100,112,125), RGBColor(38,125,103)),
        "colorful": (RGBColor(29,20,56), RGBColor(47,34,81), RGBColor(67,47,105), RGBColor(255,112,92), RGBColor(255,250,247), RGBColor(218,204,230), RGBColor(111,222,190)),
        "classic": (RGBColor(250,247,239), RGBColor(255,255,255), RGBColor(238,231,216), RGBColor(38,83,112), RGBColor(37,43,49), RGBColor(104,109,111), RGBColor(73,124,93)),
    }
    BG, PANEL, PANEL2, ACCENT, WHITE, MUTED, GREEN = palettes.get(settings.get("style", "dark"), palettes["dark"])
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    title=content.get("title") or topic; subtitle=content.get("subtitle") or "Research-based study notes"
    images=[s for s in sources if s.get("source_type")=="image"]
    slide=prs.slides.add_slide(prs.slide_layouts[6]); background(slide); rect(slide,8.65,.55,4,6.35,PANEL)
    for i,height in enumerate([1.2,2.25,1.55,3.2,2.45,3.8]): rect(slide,9.05+i*.48,5.8-height,.22,height,ACCENT if i in (1,4) else PANEL2,False)
    if images:add_image(slide,images[0],9.05,1.0,3.15,2.0)
    add_text(slide,"FIELD NOTES / RESEARCH",.82,.82,7.1,.4,11,ACCENT,True)
    add_text(slide,title,.82,1.72,7.25,2.5,34,WHITE,True)
    add_text(slide,subtitle,.86,4.55,6.8,1.0,17,MUTED)
    add_text(slide,"Explanatory study notes based on collected sources",.86,6.55,7.1,.35,10,MUTED)

    slides=content.get("slides",[])
    for idx,item in enumerate(slides):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); background(slide); rect(slide,.7,.68,.1,.62,ACCENT,False)
        add_text(slide,f"{idx+1:02d} / {len(slides):02d}",11.5,.58,1,.3,10,MUTED)
        add_text(slide,item.get("title","Key idea"),.98,.62,10.7,.85,25,WHITE,True)
        rect(slide,.85,1.65,7.75,4.95,PANEL)
        body=(item.get("body") or "\n\n".join(str(x) for x in item.get("bullets",[]))).strip()
        if not body:
            body = "\n\n".join(
                str(item.get(key, "")).strip()
                for key in ("example", "takeaway", "highlight")
                if str(item.get(key, "")).strip()
            )
        if not body:
            raise ValueError(f"Слайд {idx + 1} «{item.get('title', 'Без названия')}» пустой; PPTX не сохранён.")
        words=len(body.split()); body_size=16 if words<90 else 14.5 if words<125 else 13 if words<165 else 11.5
        add_text(slide,body,1.12,1.92,7.2,4.38,body_size,WHITE)
        rect(slide,8.9,1.65,3.55,4.95,RGBColor(17,27,45))
        img=images[(idx+1)%len(images)] if images and settings.get("visuals", "mixed") != "text" else None
        if img and add_image(slide,img,9.18,1.95,2.98,1.72):
            add_text(slide,f"Image: {img.get('license','see source page')}",9.18,3.72,3,.25,8,MUTED)
            cy=4.08
        else:
            add_text(slide,str(item.get("visual","KEY IDEA")).upper(),9.18,2.02,2.9,.3,10,GREEN,True)
            cy=2.5
        example=str(item.get("example","")).strip()
        if example:
            add_text(slide,"EXAMPLE",9.18,cy,2.9,.28,9,ACCENT,True)
            add_text(slide,example,9.18,cy+.34,2.95,1.05,11,WHITE)
            cy+=1.52
        takeaway=str(item.get("takeaway") or item.get("highlight") or "").strip()
        if takeaway:
            add_text(slide,"MAIN IDEA",9.18,cy,2.9,.28,9,GREEN,True)
            add_text(slide,takeaway,9.18,cy+.34,2.95,max(.75,6.25-cy),11,MUTED,True)
        notes=str(item.get("speaker_notes","")).strip()
        if notes:
            try: slide.notes_slide.notes_text_frame.text=notes
            except Exception: pass
        evidence=item.get("evidence",[])
        if evidence:add_text(slide,"Evidence: "+", ".join(evidence),.95,6.68,8,.22,8,MUTED)
        add_text(slide,"PRESENTATION / STUDY NOTES",.9,7.05,4.8,.2,8,MUTED)

    slide=prs.slides.add_slide(prs.slide_layouts[6]); background(slide); rect(slide,.7,.68,.1,.62,ACCENT,False)
    add_text(slide,"Sources & reading status",.98,.68,11,.8,27,WHITE,True)
    add_text(slide,"Reading status describes what the bot could extract; it does not guarantee every claim is verified.",1,1.55,11,.45,11,MUTED)
    y=2.1
    for i,source in enumerate(sources[:15],1):
        titleline=f"[S{i}] {source.get('title','Source')} — {source.get('read_status','status unknown')}"
        add_text(slide,titleline[:145],1,y,11,.28,10,WHITE,True)
        add_text(slide,source.get("url","")[:150],1,y+.27,11,.25,8,ACCENT)
        if source.get("source_type")=="image":
            add_text(slide,f"Image license: {source.get('license','see original page')}",1,y+.49,10.9,.2,8,MUTED); y+=.72
        else:y+=.57
        if y>6.8:break
    output_path=Path(output_path); output_path.parent.mkdir(parents=True,exist_ok=True); prs.save(output_path)
